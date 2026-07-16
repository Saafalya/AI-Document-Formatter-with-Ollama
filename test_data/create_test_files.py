"""
create_test_files.py
--------------------
Generates sample DOCX and PDF test files for pipeline validation.
Run: python test_data/create_test_files.py
"""

from pathlib import Path

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

BASE = Path(__file__).parent


def create_sample_docx(path: Path) -> None:
    """Create a DOCX with intentional style guide violations."""
    doc = Document()

    doc.add_heading("Introduction", level=1)
    doc.add_paragraph(
        "The quarterly report was prepared by the finance team. "
        "This document outlines the project scope and deliverables."
    )
    doc.add_paragraph(
        "Please click the submit button to finalize your entry. "
        "The chairman approved the budget after reviewing manpower requirements."
    )

    doc.add_heading("Scope", level=1)
    doc.add_paragraph(
        "We will utilize our synergy to leverage the paradigm shift. "
        "The issue was resolved by the team in a very efficient manner."
    )
    doc.add_paragraph(
        "Basically, the system is actually quite robust and really performs well. "
        "Contact support at help@acme.com or visit https://www.acme.com/docs."
    )

    doc.add_heading("Details", level=3)  # Skips H2 — hierarchy violation
    doc.add_paragraph(
        "This section discusses implementation details and technical specifications."
    )

    doc.add_heading("Conclusion", level=1)
    doc.add_paragraph(
        "In summary, the project deliverables were completed on schedule. "
        "The team recommends proceeding with phase two."
    )

    # Add a simple table
    table = doc.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Q1"
    table.rows[0].cells[2].text = "Q2"
    table.rows[0].cells[0].paragraphs[0].runs[0].bold = True
    table.rows[1].cells[0].text = "Revenue"
    table.rows[1].cells[1].text = "1,250,000"
    table.rows[1].cells[2].text = "1,380,000"

    doc.core_properties.title = "Sample Compliance Test Document"
    doc.core_properties.author = "AI Document Formatter Test Suite"

    doc.save(str(path))
    print(f"Created: {path}")


def create_style_guide_pdf(path: Path) -> None:
    """Create a PDF style guide with extractable rules."""
    c = canvas.Canvas(str(path), pagesize=letter)
    width, height = letter
    y = height - 72

    lines = [
        "ACME CORPORATION — CORPORATE STYLE GUIDE",
        "",
        "1. FORMATTING REQUIREMENTS",
        "Font: Calibri, 11pt for body text.",
        "Heading 1: 16pt bold. Heading 2: 14pt bold. Heading 3: 12pt bold.",
        "Alignment: Left-aligned body text.",
        "Margins: 1.25 inches left and right, 1 inch top and bottom.",
        "Do not use ALL CAPS for emphasis. Avoid underlining except hyperlinks.",
        "",
        "2. WRITING REQUIREMENTS",
        "Use active voice instead of passive voice.",
        "Do not begin instructions with 'Please'. Use direct imperatives.",
        "Use inclusive language throughout all documents.",
        "Avoid filler words: very, basically, actually, really, just.",
        "Preferred terminology: use 'use' instead of 'utilize' or 'leverage'.",
        "",
        "3. INCLUSIVE LANGUAGE",
        "chairman -> chairperson",
        "manpower -> workforce",
        "",
        "4. STRUCTURE REQUIREMENTS",
        "Maintain heading hierarchy. Do not skip heading levels.",
        "Required sections: Introduction, Scope, Conclusion.",
        "Maximum heading depth: 3 levels.",
        "",
        "5. TABLE FORMATTING",
        "Table header rows must be bold.",
        "",
        "6. BRANDING",
        "Company name: Acme Corporation",
        "Primary color: #003366",
    ]

    c.setFont("Helvetica-Bold", 14)
    c.drawString(72, y, lines[0])
    y -= 30
    c.setFont("Helvetica", 10)

    for line in lines[2:]:
        if y < 72:
            c.showPage()
            y = height - 72
            c.setFont("Helvetica", 10)
        c.drawString(72, y, line)
        y -= 14

    c.save()
    print(f"Created: {path}")


if __name__ == "__main__":
    create_sample_docx(BASE / "sample_document.docx")
    create_style_guide_pdf(BASE / "style_guide.pdf")
    print("Test files ready in test_data/")
